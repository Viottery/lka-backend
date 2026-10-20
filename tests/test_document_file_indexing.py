from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.api.routes.knowledge import router
from app.domains.file_processing import FileProcessingError, FileProcessor, file_type_policies
from app.domains.knowledge import KnowledgeService
from app.domains.workspace_knowledge import WorkspaceKnowledgeIndexer
from app.storage.db import connect, get_db_path, init_db


def _service(tmp_path: Path, root: Path) -> KnowledgeService:
    db_path = get_db_path(tmp_path / "data")
    init_db(db_path)
    return KnowledgeService(lambda: connect(db_path), workspace_roots=(root,))


def _pdf(path: Path, text: str = "") -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def _docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>wordunique evidence</w:t></w:r></w:p></w:body></w:document>')


def test_workspace_converts_mixed_documents_and_preserves_evidence_metadata(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    (root / "notes.md").write_text("nativeunique evidence", encoding="utf-8")
    (root / "table.CSV").write_text("name,value\ncsvunique,37\n", encoding="utf-8")
    (root / "data.json").write_text('{"value":"jsonunique"}', encoding="utf-8")
    (root / "page.html").write_text("<h1>htmlunique evidence</h1>", encoding="utf-8")
    _pdf(root / "report.pdf", "pdfunique evidence")
    _docx(root / "word.docx")
    book = Workbook()
    book.active.append(["Name", "Value"])
    book.active.append(["sheetunique", 48])
    book.save(root / "sheet.xlsx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "slideunique evidence"
    slides.save(root / "slides.pptx")
    (root / "unknown.bin").write_bytes(b"ignoredunique evidence")
    service = _service(tmp_path, root)

    result = WorkspaceKnowledgeIndexer(service, [root]).index_workspace(root)

    assert result.imported_files == result.discovered_files == 8
    assert result.failed_files == 0
    assert not result.errors
    for keyword, name, processor in [
        ("nativeunique", "notes.md", "native_text"),
        ("csvunique", "table.CSV", "markitdown"),
        ("jsonunique", "data.json", "markitdown"),
        ("htmlunique", "page.html", "markitdown"),
        ("pdfunique", "report.pdf", "native_pdf"),
        ("wordunique", "word.docx", "markitdown"),
        ("sheetunique", "sheet.xlsx", "markitdown"),
        ("slideunique", "slides.pptx", "markitdown"),
    ]:
        hit = service.search(query=keyword, mode="keyword").results[0]
        extraction = hit.metadata["extraction"]
        assert extraction["processor"] == processor
        assert extraction["source_sha256"] == hashlib.sha256((root / name).read_bytes()).hexdigest()
        assert extraction["fidelity_verified"] is False
        chunk = service.load_chunks(chunk_ids=[hit.chunk_id]).chunks[0]
        assert chunk.metadata == hit.metadata
        document = service.load_document(document_id=hit.document_id)
        assert document.metadata["relative_path"] == name
        assert document.metadata["extraction"]["processor"] == processor
        assert hit.source_ref.startswith("workspace_file:workspace://")
        if name == "report.pdf":
            assert extraction["pages"] == [1]
            assert extraction["warnings"]
            assert document.metadata["extraction"]["page_count"] == 1
    assert not service.search(query="ignoredunique", mode="keyword").results
    assert not service.list_authorized_source_ids()
    assert service.list_authorized_source_ids(workspace_path=str(root))


def test_failed_replacement_retains_old_document_and_does_not_prune(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    pdf = root / "report.pdf"
    _pdf(pdf, "oldpdfunique evidence")
    note = root / "note.txt"
    note.write_text("oldnoteunique evidence", encoding="utf-8")
    service = _service(tmp_path, root)
    indexer = WorkspaceKnowledgeIndexer(service, [root])
    assert indexer.index_workspace(root).imported_files == 2
    original = service.search(query="oldpdfunique", mode="keyword").results[0]
    _pdf(pdf)
    note.unlink()

    failed = indexer.index_workspace(root)

    assert failed.failed_files == 1
    assert failed.pruned_documents == 0
    assert service.search(query="oldpdfunique", mode="keyword").results[0].document_id == original.document_id
    assert service.search(query="oldnoteunique", mode="keyword").results
    assert str(root) not in json.dumps(failed.errors)
    _pdf(pdf, "newpdfunique evidence")
    success = indexer.index_workspace(root)
    assert success.imported_files == 1
    assert success.pruned_documents == 1
    assert not service.search(query="oldpdfunique", mode="keyword").results
    assert not service.search(query="oldnoteunique", mode="keyword").results
    assert service.search(query="newpdfunique", mode="keyword").results


def test_over_limit_file_does_not_remove_previous_index(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    note = root / "note.txt"
    note.write_text("retainedunique evidence", encoding="utf-8")
    service = _service(tmp_path, root)
    assert WorkspaceKnowledgeIndexer(service, [root]).index_workspace(root).imported_files == 1
    result = WorkspaceKnowledgeIndexer(service, [root], max_file_bytes=2).index_workspace(root)
    assert result.skipped_files == 1
    assert result.errors
    assert result.pruned_documents == 0
    assert service.search(query="retainedunique", mode="keyword").results


def test_corrupt_office_cannot_fall_back_to_zip_and_replace_evidence(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    path = root / "report.docx"
    _docx(path)
    service = _service(tmp_path, root)
    indexer = WorkspaceKnowledgeIndexer(service, [root])
    assert indexer.index_workspace(root).imported_files == 1
    before = service.search(query="wordunique", mode="keyword").results[0]
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("unrelated.txt", "fallbackunique false evidence")

    result = indexer.index_workspace(root)

    assert result.failed_files == 1
    assert result.pruned_documents == 0
    assert service.search(query="wordunique", mode="keyword").results[0].document_id == before.document_id
    assert not service.search(query="fallbackunique", mode="keyword").results


def test_single_file_import_uses_snapshot_and_rejects_empty_replacement(tmp_path, monkeypatch):
    root = tmp_path / "workspace"
    root.mkdir()
    path = root / "table.csv"
    original = b"name,value\nsnapshotunique,37\n"
    path.write_bytes(original)
    service = _service(tmp_path, root)
    process = service.file_processor.process_bytes

    def remove_after_read(filename, raw):
        assert filename == path.name and raw == original
        path.unlink()
        return process(filename, raw)

    monkeypatch.setattr(service.file_processor, "process_bytes", remove_after_read)
    imported = service.import_file(path=path)
    assert service.search(query="snapshotunique", mode="keyword").results
    monkeypatch.setattr(service.file_processor, "process_bytes", process)
    path.write_bytes(b"")
    with pytest.raises(FileProcessingError):
        service.import_file(path=path)
    assert service.load_document(document_id=imported.document_id).chunk_ids
    outside = tmp_path / "outside.csv"
    outside.write_bytes(original)
    with pytest.raises(ValueError, match="outside configured workspace roots"):
        service.import_file(path=outside)


def test_file_types_endpoint_matches_processing_registry():
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as client:
        response = client.get("/knowledge/file-types")
    assert response.status_code == 200
    assert response.json()["file_types"] == file_type_policies()
    policies = {extension: row for row in response.json()["file_types"] for extension in row["extensions"]}
    assert policies[".pdf"]["preferred_processor"] == "native_pdf"
    assert policies[".docx"]["preferred_processor"] == "markitdown"
    assert policies[".png"]["index_supported"] is False
    assert policies[".zip"]["index_supported"] is False


def test_pdf_provenance_can_locate_chunks_after_page_splits(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    path = root / "large.pdf"
    writer = PdfWriter()
    for index in range(2):
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({'longunique ' * 240} PAGE{index+1}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    path.write_bytes(buffer.getvalue())
    service = _service(tmp_path, root)
    imported = service.import_file(path=path)
    doc = service.load_document(document_id=imported.document_id)
    text = FileProcessor().process_bytes(path.name, path.read_bytes()).text
    chunks = service.load_chunks(chunk_ids=doc.chunk_ids, max_chars_per_chunk=6000).chunks
    assert len(chunks) > 2
    for chunk in chunks:
        extraction = chunk.metadata["extraction"]
        start, end = extraction["source_start_char"], extraction["source_end_char"]
        assert text[start:end] == chunk.text
        assert extraction["pages"]
