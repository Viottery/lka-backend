from __future__ import annotations

import hashlib
import io
import zipfile

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.domains.file_processing import (
    FileProcessingError,
    FileProcessor,
    file_type_policies,
    get_file_type_policy,
    supported_index_extensions,
)


def _zip(members):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, text in members.items():
            archive.writestr(name, text)
    return stream.getvalue()


def _pdf(text="PDF sentinel"):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def test_registry_explicit_and_json_safe():
    assert get_file_type_policy("REPORT.PDF").preferred_processor == "native_pdf"
    assert get_file_type_policy("image.png").index_supported is False
    assert get_file_type_policy("file.unknown") is None
    assert {".pdf", ".xlsx", ".md"} <= supported_index_extensions()
    assert ".zip" not in supported_index_extensions()
    assert all(isinstance(item["extensions"], list) for item in file_type_policies())


def test_native_text_snapshot_and_provenance():
    raw = b"\xef\xbb\xbf  snapshot sentinel \n"
    result = FileProcessor().process_bytes("unavailable/path/note.TXT", raw)
    assert result.text == "snapshot sentinel"
    assert result.mime_type == "text/plain"
    assert result.metadata["extraction"]["source_sha256"] == hashlib.sha256(raw).hexdigest()


@pytest.mark.parametrize("extension,raw", [
    (".csv", b"name,value\nsentinel,42\n"),
    (".html", b"<html><body><h1>sentinel</h1><p>local content</p></body></html>"),
    (".json", b'{"sentinel": 42}'),
    (".xml", b"<root><sentinel>42</sentinel></root>"),
])
def test_real_text_conversions(extension, raw):
    result = FileProcessor().process_bytes("absent" + extension, raw)
    assert "sentinel" in result.text
    assert result.metadata["extraction"]["processor"] == "markitdown"


def test_attachment_processing_can_raise_original_size_limit(monkeypatch):
    monkeypatch.setattr("app.domains.file_processing.MAX_SOURCE_BYTES", 10)
    processor = FileProcessor()
    with pytest.raises(FileProcessingError):
        processor.process_bytes("note.txt", b"x" * 11)
    result = processor.process_bytes("note.txt", b"x" * 11, source_max_bytes=16)
    assert result.text == "x" * 11


def test_real_docx_conversion():
    raw = _zip({
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>DOCX sentinel</w:t></w:r></w:p></w:body></w:document>',
    })
    assert "DOCX sentinel" in FileProcessor().process_bytes("source.docx", raw).text


def test_real_pptx_conversion():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX sentinel"
    stream = io.BytesIO()
    presentation.save(stream)
    assert "PPTX sentinel" in FileProcessor().process_bytes("source.pptx", stream.getvalue()).text


def test_real_xlsx_conversion():
    workbook = Workbook()
    workbook.active.append(["XLSX sentinel", 42])
    stream = io.BytesIO()
    workbook.save(stream)
    assert "XLSX sentinel" in FileProcessor().process_bytes("source.xlsx", stream.getvalue()).text


def test_pdf_snapshot_has_page_spans_and_visual_limitations():
    result = FileProcessor().process_bytes("missing.pdf", _pdf())
    extraction = result.metadata["extraction"]
    assert extraction["processor"] == "native_pdf"
    assert extraction["text_layer_only"] is True
    assert extraction["fidelity_verified"] is False
    span = extraction["page_spans"][0]
    assert result.text[span["start_char"]:span["end_char"]] == "PDF sentinel"
    assert extraction["warnings"]


@pytest.mark.parametrize("filename,raw", [("scan.pdf", _pdf("")), ("empty.txt", b" \n"),
                                         ("bad.pdf", b"private contents"),
                                         ("bad.docx", b"private contents"),
                                         ("binary.png", b"sentinel")])
def test_failure_does_not_emit_empty_or_binary_fallback(filename, raw):
    with pytest.raises(FileProcessingError) as failure:
        FileProcessor().process_bytes(filename, raw)
    assert "private contents" not in str(failure.value)


def test_archive_member_limit(monkeypatch):
    monkeypatch.setattr("app.integrations.document_conversion.MAX_ARCHIVE_MEMBERS", 1)
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("source.docx", _zip({"one": "x", "two": "x"}))


def test_archive_expanded_byte_limit(monkeypatch):
    monkeypatch.setattr("app.integrations.document_conversion.MAX_ARCHIVE_BYTES", 10)
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("source.xlsx", _zip({"one": "x" * 11}))


def test_input_and_output_limits(monkeypatch):
    monkeypatch.setattr("app.domains.file_processing.MAX_SOURCE_BYTES", 10)
    with pytest.raises(FileProcessingError, match="processing limit"):
        FileProcessor().process_bytes("large.txt", b"x" * 11)
    monkeypatch.setattr("app.domains.file_processing.MAX_EXTRACTED_CHARS", 2)
    with pytest.raises(FileProcessingError, match="character limit"):
        FileProcessor().process_bytes("note.txt", b"abc")


def test_pdf_never_returns_partial_index(monkeypatch):
    with pytest.raises(ValueError, match="character limit"):
        FileProcessor()._documents.extract_snapshot(_pdf(), max_chars=5)


def test_converter_disables_plugins_and_uses_only_bytes(monkeypatch):
    import markitdown
    observed = {}
    class LocalConverter:
        def __init__(self, **kwargs):
            observed["options"] = kwargs
        def register_converter(self, converter):
            observed["registered"] = type(converter).__name__
        def convert_stream(self, source, **kwargs):
            observed["snapshot"] = source.read()
            observed["conversion"] = kwargs
            return type("Result", (), {"text_content": "local sentinel"})()
    monkeypatch.setattr(markitdown, "MarkItDown", LocalConverter)
    result = FileProcessor().process_bytes("https://private.invalid/source.html", b"snapshot")
    assert result.text == "local sentinel"
    assert observed["options"] == {"enable_builtins": False, "enable_plugins": False}
    assert observed["registered"] == "HtmlConverter"
    assert observed["snapshot"] == b"snapshot"
    assert observed["conversion"]["keep_data_uris"] is False


@pytest.mark.parametrize("extension", ["pptx", "xlsx"])
def test_empty_office_scaffolding_is_not_content(extension):
    document = Presentation() if extension == "pptx" else Workbook()
    if extension == "pptx":
        document.slides.add_slide(document.slide_layouts[6])
    stream = io.BytesIO()
    document.save(stream)
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("empty." + extension, stream.getvalue())


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "epub"])
def test_malformed_archive_never_falls_back_to_member_import(extension):
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("broken." + extension, _zip({"secret.txt": "sentinel"}))


@pytest.mark.parametrize("extension", ["xls", "msg"])
def test_binary_document_never_falls_back_to_plain_text(extension):
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("broken." + extension, b"plain sentinel")


@pytest.mark.parametrize("extension", ["html", "csv", "json", "xml"])
def test_binary_source_disguised_as_text_is_rejected(extension):
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("broken." + extension, _zip({"secret.txt": "sentinel"}))


def _epub(chapter=None, reference="chapter.xhtml"):
    members = {
        "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>',
        "OPS/book.opf": f'<package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>CompleteBook</dc:title></metadata><manifest><item id="chapter" href="{reference}" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>',
    }
    if chapter is not None:
        members["OPS/chapter.xhtml"] = chapter
    return _zip(members)


def test_epub_preserves_actual_chapter_text():
    raw = _epub('<html xmlns="http://www.w3.org/1999/xhtml"><body><p>chapter sentinel</p></body></html>')
    assert "chapter sentinel" in FileProcessor().process_bytes("valid.epub", raw).text


@pytest.mark.parametrize("raw", [
    _epub(),
    _epub('<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Only title</title></head><body/></html>'),
    _epub('<html xmlns="http://www.w3.org/1999/xhtml"><body><img src="cover.jpg" alt="cover"/></body></html>'),
    _epub(reference="https://private.invalid/chapter.xhtml"),
])
def test_epub_missing_or_empty_chapters_do_not_index_metadata(raw):
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("incomplete.epub", raw)


def test_image_placeholders_do_not_qualify_as_content():
    from app.integrations.document_conversion import _meaningful_text
    assert not _meaningful_text('<!-- Slide number: 1 -->\n![cover](data:image/png;base64,aaa)', '.pptx')
    assert _meaningful_text('![cover](image.png)\nactual text', '.docx') == 'actual text'


def test_image_only_pptx_is_not_extractable_text():
    from PIL import Image
    from pptx.util import Inches
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "red").save(image, format="PNG")
    image.seek(0)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(image, Inches(1), Inches(1))
    stream = io.BytesIO()
    presentation.save(stream)
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("image-only.pptx", stream.getvalue())


def test_image_only_docx_is_not_extractable_text():
    from PIL import Image
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "red").save(image, format="PNG")
    raw = _zip({
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image.png"/></Relationships>',
        "word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body><w:p><w:r><w:drawing><wp:inline><wp:docPr id="1" name="Picture"/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:blipFill><a:blip r:embed="rId1"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p></w:body></w:document>',
        "word/media/image.png": image.getvalue(),
    })
    with pytest.raises(FileProcessingError):
        FileProcessor().process_bytes("image-only.docx", raw)
